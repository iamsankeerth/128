#!/usr/bin/env python3
"""Regenerate JSON data from GMAT Excel workbooks."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"


def parse_range(value):
    if not value:
        return []
    text = str(value).strip()
    if "–" in text:
        start, end = text.split("–", 1)
        return list(range(int(start), int(end) + 1))
    if "-" in text and text[0].isdigit():
        start, end = text.split("-", 1)
        return list(range(int(start), int(end) + 1))
    return [int(part) for part in re.split(r"[,\s]+", text) if part.isdigit()]


def load_questions(master_path: Path):
    workbook = load_workbook(master_path, data_only=True)
    sheet = workbook["Master Index"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    questions = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        record = dict(zip(headers, row))
        number = int(record["Question No."])
        questions[number] = {
            "no": number,
            "id": record.get("Question ID"),
            "section": record.get("Section"),
            "concept": record.get("Official Concept"),
            "subtype": record.get("Detailed Subtype"),
            "difficulty": record.get("Difficulty"),
            "bookPage": record.get("Book Q Page"),
            "pdfPage": record.get("PDF Q Page"),
            "explanationPage": record.get("PDF Explanation Page"),
            "passageSet": record.get("Passage Set"),
            "correctAnswer": record.get("Correct Answer"),
        }
    return questions


def load_daily_plan(daily_path: Path):
    workbook = load_workbook(daily_path, data_only=True)
    sheet = workbook.active
    days = []
    for row in sheet.iter_rows(values_only=True):
        if not row or not row[0]:
            continue
        first = str(row[0]).strip()
        if first in {"Daily Division", "Date", "Plan Summary", "MONTH TOTAL"} or first.startswith("The "):
            continue
        if first.startswith(("Study Days", "RC ", "CR ", "Total Questions", "Average")):
            continue
        date_value = row[0]
        if isinstance(date_value, str) and not re.search(r"\d", date_value):
            continue
        rc_numbers = parse_range(row[2])
        cr_numbers = parse_range(row[4])
        if not rc_numbers and not cr_numbers:
            continue
        if hasattr(date_value, "strftime"):
            date_iso = date_value.strftime("%Y-%m-%d")
            label = date_value.strftime("%a, %d-%b")
        else:
            date_iso = str(date_value)
            label = str(date_value)
        days.append(
            {
                "date": date_iso,
                "label": label,
                "passageSets": row[1],
                "rcRange": str(row[2] or ""),
                "rcCount": int(row[3] or len(rc_numbers)),
                "crRange": str(row[4] or ""),
                "crCount": int(row[5] or len(cr_numbers)),
                "total": int(row[6] or (len(rc_numbers) + len(cr_numbers))),
                "rcQuestions": rc_numbers,
                "crQuestions": cr_numbers,
                "questionNumbers": rc_numbers + cr_numbers,
            }
        )
    return {"days": days, "summary": {"totalQuestions": 346, "studyDays": len(days)}}


def main():
    daily_path = Path("/tmp/daily.xlsx")
    master_path = Path("/tmp/master.xlsx")
    if not daily_path.exists() or not master_path.exists():
        raise SystemExit("Place daily.xlsx and master.xlsx in /tmp before running.")

    DATA.mkdir(parents=True, exist_ok=True)
    (DATA / "questions.json").write_text(json.dumps(load_questions(master_path), indent=2))
    (DATA / "daily-plan.json").write_text(json.dumps(load_daily_plan(daily_path), indent=2))
    print("Wrote", DATA / "questions.json")
    print("Wrote", DATA / "daily-plan.json")


if __name__ == "__main__":
    main()
